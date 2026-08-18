from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "competition" / "真信盾项目功能与技术说明书（完善版）.docx"
POSTER = ROOT / "docs" / "competition" / "ada1f1106d5ee48438fd6e1d596e200e.jpg"

NAVY = "0B1F3A"
BLUE = "1677FF"
CYAN = "00B8D9"
LIGHT = "EAF4FF"
PALE = "F4F7FB"
GRAY = "5E6B7A"
WHITE = "FFFFFF"
GOLD = "D69E2E"
RED = "B42318"
GREEN = "067647"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.78)
sec.bottom_margin = Inches(0.72)
sec.left_margin = Inches(0.82)
sec.right_margin = Inches(0.82)
sec.header_distance = Inches(0.35)
sec.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(NAVY)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in [
    ("Title", 28, NAVY, 0, 8),
    ("Heading 1", 18, BLUE, 16, 8),
    ("Heading 2", 14, NAVY, 12, 6),
    ("Heading 3", 11.5, "245A8D", 9, 4),
]:
    st = styles[name]
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for custom, size, color, bold in [
    ("Kicker", 9, CYAN, True),
    ("Lead", 12, NAVY, False),
    ("Small", 8.5, GRAY, False),
]:
    st = styles.add_style(custom, WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = "Microsoft YaHei"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = bold

def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")

def set_table_widths(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW"); tbl_pr.append(tbl_w)
    total = sum(widths)
    tbl_w.set(qn("w:w"), str(total)); tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd"); tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120"); tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(width)); grid.append(gc)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width)); tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def style_table(table, header=True, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    if widths: set_table_widths(table, widths)
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            if header and ri == 0: set_cell_fill(cell, NAVY)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.12
                for run in p.runs:
                    run.font.name = "Microsoft YaHei"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                    run.font.size = Pt(9)
                    if header and ri == 0:
                        run.font.bold = True; run.font.color.rgb = RGBColor.from_string(WHITE)

def add_table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers): t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row): cells[i].text = str(value)
    style_table(t, True, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def add_bullets(items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25 + 0.2 * level)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)

def create_numbering_id():
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids or [0]) + 1
    num_id = max(num_ids or [0]) + 1
    abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid"); nsid.set(qn("w:val"), f"{abstract_id:08X}"); abstract.append(nsid)
    multi = OxmlElement("w:multiLevelType"); multi.set(qn("w:val"), "singleLevel"); abstract.append(multi)
    lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
    num_fmt = OxmlElement("w:numFmt"); num_fmt.set(qn("w:val"), "decimal"); lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText"); lvl_text.set(qn("w:val"), "%1."); lvl.append(lvl_text)
    suff = OxmlElement("w:suff"); suff.set(qn("w:val"), "tab"); lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), "540"); tabs.append(tab); p_pr.append(tabs)
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "540"); ind.set(qn("w:hanging"), "270"); p_pr.append(ind)
    lvl.append(p_pr); abstract.append(lvl); numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId"); abstract_ref.set(qn("w:val"), str(abstract_id)); num.append(abstract_ref); numbering.append(num)
    lvl_override = OxmlElement("w:lvlOverride"); lvl_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride"); start_override.set(qn("w:val"), "1")
    lvl_override.append(start_override); num.append(lvl_override)
    return num_id

def add_numbered(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)

def callout(label, text, fill=LIGHT, color=NAVY):
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0); set_cell_fill(c, fill); set_cell_margins(c, 150, 180, 150, 180)
    p = c.paragraphs[0]
    r = p.add_run(label + "  "); r.bold = True; r.font.color.rgb = RGBColor.from_string(BLUE)
    r = p.add_run(text); r.font.color.rgb = RGBColor.from_string(color)
    set_table_widths(t, [9360])
    t.style = "Table Grid"
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def add_page_number(paragraph):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)

# Header/footer
header = sec.header
hp = header.paragraphs[0]
hp.text = "真信盾 True Shield  |  项目功能与技术说明书"
hp.style = styles["Small"]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("内部项目资料  ·  2026-08-09  ·  第 ")
run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(GRAY)
add_page_number(fp)
fp.add_run(" 页").font.size = Pt(8)

# Cover
p = doc.add_paragraph(style="Kicker"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("FAMILY-GRADE MULTIMODAL ANTI-FRAUD SYSTEM")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("真信盾"); r.bold = True; r.font.name = "Microsoft YaHei"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei"); r.font.size = Pt(34); r.font.color.rgb = RGBColor.from_string(NAVY)
p.paragraph_format.space_after = Pt(3)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("家庭级多模态智能反诈与可信联系系统"); r.font.size = Pt(16); r.font.color.rgb = RGBColor.from_string(BLUE)
p.paragraph_format.space_after = Pt(16)
if POSTER.exists():
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(POSTER), width=Inches(6.65))
    p.paragraph_format.space_after = Pt(16)
callout("项目定位", "在转账、泄密、点击可疑链接等高风险操作发生前，通过多模态风险识别、可解释证据和家庭协同干预，帮助用户停下来、看证据、找家人、再核验。")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("项目功能、技术架构、安全设计与实施边界综合说明"); r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string(GRAY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("依据当前代码仓库整理  |  版本快照：2026年8月9日"); r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(GRAY)
doc.add_page_break()

# Contents
doc.add_heading("文档导读", level=1)
doc.add_paragraph("本说明书以当前仓库代码、依赖配置、数据库迁移与项目文档为依据，重点回答“已经实现什么、如何实现、如何演示、哪些仍是边界”。")
add_table(["章节", "内容"], [
    ("1-2", "项目概述、用户痛点与价值定位"),
    ("3", "已实现功能全景与状态"),
    ("4-5", "系统架构、技术栈与分析引擎"),
    ("6-8", "家庭协同、通话护航、数据与接口"),
    ("9-11", "安全隐私、测试质量、部署运行"),
    ("12-16", "创新点、比赛演示、边界、持续完善计划与路线图"),
], [1440, 7920])
callout("状态口径", "“已实现”表示仓库存在完整业务链路；“可选配置”表示代码已具备但依赖外部凭证、模型或通道；“待生产化”表示原型可用，但尚缺少生产级运营、监控、限流或真实规模验证。", PALE)

doc.add_heading("1. 项目概述", level=1)
doc.add_paragraph("真信盾面向家庭反诈场景，服务对象包括老年人、青少年、数字技能较弱用户及其家属。系统通过Android客户端采集用户主动提交的文本、截图、语音、网址或号码信息，由FastAPI后端与本地规则共同分析，输出风险等级、命中证据和行动建议；高风险事件可进一步进入家庭告警与可信联系人协同流程。")
add_table(["维度", "说明"], [
    ("核心问题", "诈骗跨文本、图片、语音、网址和电话渠道传播，个人在紧迫、恐惧或权威压力下容易独自误判。"),
    ("核心方案", "多模态识别 + 可解释规则/AI融合 + 家庭可信联系人 + 风险记录与复盘。"),
    ("产品形态", "Android应用、FastAPI服务端、PostgreSQL数据层及Docker Compose本地基础设施。"),
    ("结果定位", "提供风险提示和核验建议，不替代公安、银行、运营商或司法机构作出权威结论。"),
], [1800, 7560])

doc.add_heading("2. 用户场景与价值闭环", level=1)
add_numbered([
    "用户收到可疑聊天、短信截图、语音、网址或陌生来电。",
    "用户主动提交内容，或来电护航在系统回调中读取号码及验证状态。",
    "系统进行规则识别、OCR/ASR转换、URL/号码信号分析，并可选择AI语义复核。",
    "客户端展示风险等级、证据、解释与建议，避免只给出黑箱分数。",
    "高风险事件可触发家庭告警，可信联系人通过推送或邮件获知并协助核验。",
    "检测结果进入历史、详情、仪表盘、反馈与PDF报告，形成复盘闭环。",
])

doc.add_heading("3. 已实现功能全景", level=1)
add_table(["功能域", "已实现能力", "状态"], [
    ("账户与安全", "注册、登录、JWT鉴权、个人资料、修改密码、验证码重置、旧Token失效", "已实现"),
    ("文本检测", "规则匹配、组合条件、上下文/否定抑制、证据、分数、等级、建议", "已实现"),
    ("图片检测", "图片上传、RapidOCR文字提取、风险分析、OCR纠正与重新分析", "已实现"),
    ("语音检测", "音频上传、格式/大小/时长校验、FunASR/SenseVoice转写、风险分析", "可选配置"),
    ("网址检测", "URL结构、重定向与风险信号分析，输出证据和行动建议", "已实现"),
    ("号码与来电", "号码筛查、本地离线规则、CallScreeningService、提醒、可选静音/确认风险拦截", "已实现"),
    ("风险事件", "事件保存、信号证据、历史筛选、详情、统计仪表盘、PDF报告", "已实现"),
    ("反馈纠错", "正确/误报/漏报/等级偏差反馈，OCR纠正，重新分析，反馈统计", "已实现"),
    ("家庭体系", "家庭创建、邀请、成员与角色、可信联系人、通知偏好", "已实现"),
    ("家庭告警", "手动/策略自动创建、接收人选择、投递记录、失败原因、重试与冷却", "已实现"),
    ("消息通知", "FCM设备注册与推送、邮件通知、投递状态记录", "可选配置"),
    ("报告与复盘", "风险看板、来源/等级趋势、纠错统计、PDF风险报告", "已实现"),
], [1650, 6200, 1510])

doc.add_heading("3.1 账户、家庭与可信联系人", level=2)
add_bullets([
    "账户生命周期覆盖注册、登录、资料维护、修改密码和密码重置。",
    "家庭资源访问要求有效成员身份与角色权限，避免通过UUID枚举访问他人数据。",
    "家庭邀请、成员管理、可信联系人及其邮件/推送偏好形成协作关系基础。",
    "密码修改后可使旧Token及推送设备认证状态失效，减少凭证长期暴露风险。",
])

doc.add_heading("3.2 多模态风险检测", level=2)
add_bullets([
    "文本：直接进入规则引擎，可选AI语义复核。",
    "图片：先执行OCR，再将提取文本送入统一风险分析；支持用户纠正OCR结果。",
    "语音：先执行本地/服务端ASR转写，再分析诱导转账、索要验证码、冒充身份和紧急施压等语义。",
    "网址：分析协议、域名、可疑结构、跳转链和其他链接信号。",
    "电话：本地规则在来电响应窗口内快速判断，网络结果作为后续补充，不阻塞系统回调。",
])

doc.add_page_break()
doc.add_heading("4. 系统架构", level=1)
callout("当前架构", "项目采用模块化单体，而非已经拆分完成的微服务。services与packages中的部分目录属于未来拆分预留，不应在材料中宣传为已部署微服务。", "FFF4E5", "6B4F00")
p = doc.add_paragraph()
p.style = styles["Lead"]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Android Compose  →  ViewModel / StateFlow  →  Repository  →  Retrofit / Bearer Token\n")
p.add_run("FastAPI Route  →  Domain Service / Analyzer  →  SQLAlchemy Repository  →  PostgreSQL")
add_table(["层次", "职责"], [
    ("Android UI", "Jetpack Compose页面、Material 3组件、导航、权限引导和状态展示。"),
    ("状态与数据", "ViewModel、StateFlow、Repository负责页面状态、校验与网络操作。"),
    ("API层", "FastAPI路由、Pydantic Schema、JWT依赖与OpenAPI文档。"),
    ("领域服务", "规则引擎、AI融合、OCR、ASR、URL、号码、报告、通知和告警编排。"),
    ("持久化", "异步SQLAlchemy模型与Repository，Alembic管理迁移，PostgreSQL为生产行为基准。"),
    ("基础设施", "Docker Compose编排API、PostgreSQL、Redis和MinIO；当前核心链路主要依赖PostgreSQL。"),
], [1800, 7560])

doc.add_heading("5. 技术栈", level=1)
add_table(["领域", "技术/版本", "用途"], [
    ("Android", "Kotlin 2.0.21；JVM 11；minSdk 26；target/compile 36", "移动端开发与系统能力接入"),
    ("界面", "Jetpack Compose；Material 3；Navigation Compose 2.9.8", "声明式UI、页面导航与设计系统"),
    ("状态", "ViewModel；Lifecycle；StateFlow", "单向状态管理和生命周期感知"),
    ("网络", "Retrofit 3.0.0；Gson；OkHttp 4.12.0", "REST API调用、序列化和调试日志"),
    ("推送", "Firebase BoM 34.16.0；Firebase Messaging", "家庭告警与设备推送"),
    ("后端", "Python 3.12；FastAPI；Pydantic Settings", "异步API、校验、配置和OpenAPI"),
    ("数据", "SQLAlchemy Async；Alembic；asyncpg；PostgreSQL 16", "持久化、迁移和异步数据库访问"),
    ("OCR", "RapidOCR 3.9.2；Pillow 12.3+", "截图与图片文字识别"),
    ("语音", "FunASR 1.4；ModelScope；SenseVoice；PyTorch/Torchaudio；ONNX Runtime", "语音转写及本地模型推理"),
    ("AI语义", "可选Ollama兼容本地语义模型", "对规则结果进行语义复核与融合"),
    ("报告", "ReportLab 4.4+", "生成风险PDF报告"),
    ("安全", "Argon2（pwdlib）；PyJWT；HMAC-SHA256", "密码摘要、Token与号码反馈指纹"),
    ("基础设施", "Docker Compose；Redis 7.4；MinIO", "本地编排、缓存/对象存储预留与扩展"),
    ("质量", "pytest、pytest-asyncio、pytest-cov、Ruff、mypy、JUnit", "单元、集成、慢速测试与静态检查"),
], [1500, 3500, 4360])

doc.add_heading("5.1 规则与AI融合分析", level=2)
doc.add_paragraph("规则引擎使用YAML维护可解释规则，支持关键词、正则、组合条件、上下文、排除条件及否定语义抑制。启用AI时，语义结果与规则结果按置信度和安全边界融合，规则证据仍被保留，用于解释、复核与审计。AI不可用时，系统仍可使用规则引擎完成基础分析。")
add_bullets([
    "优势：规则可审计、可版本化、结果稳定；AI能够发现隐蔽意图和跨句行为链。",
    "安全边界：规则零分不等同于安全；AI输出不能直接覆盖明确高危证据。",
    "输出结构：风险等级、分数、证据列表、命中词/位置、解释、建议动作及规则版本。",
])

doc.add_heading("6. 家庭协同与通知闭环", level=1)
add_numbered([
    "用户创建家庭并邀请成员，配置可信联系人及通知偏好。",
    "风险事件达到家庭策略阈值时自动创建告警，也可由用户手动发起。",
    "告警服务筛选符合条件的接收人，并调用已配置的FCM或邮件提供方。",
    "系统记录每次投递渠道、状态、失败原因、尝试次数和冷却时间。",
    "家属在消息中心或告警详情中查看风险证据，联系用户并更新处置状态。",
])
callout("核心差异化", "传统检测工具往往只弹出个人警告；真信盾把“个人识别”延伸为“可信家属参与核验”，尤其适合老年人与数字弱势群体。")

doc.add_heading("7. 通话护航功能", level=1)
add_table(["能力", "实现方式", "关键边界"], [
    ("系统来电接入", "Android CallScreeningService读取号码、方向及系统号码验证状态", "需要用户授予来电筛查角色"),
    ("离线筛查", "登录后同步并缓存版本化规则包；来电回调仅读取未过期本地缓存", "不等待网络，满足系统响应窗口"),
    ("风险提示", "结合隐藏号码、验证状态、陌生来电及精确/号段规则给出等级和证据", "号码筛查不能代替人工核验"),
    ("静音", "用户可主动开启高风险来电静音", "默认仅提醒"),
    ("拦截", "仅对服务端标记为已复核、confirmed的精确高风险号码生效", "不因陌生、号段或待审核举报自动拦截"),
    ("号码反馈", "提交可疑、诈骗或误报；保存HMAC指纹与脱敏号码", "反馈进入待复核队列，不直接改变规则"),
    ("家庭求助", "通话期间可创建高风险事件并通知家庭联系人", "保存脱敏号码，不保存实时通话音频"),
], [1500, 5100, 2760])

doc.add_heading("8. 数据模型与API", level=1)
doc.add_heading("8.1 核心数据实体", level=2)
add_bullets([
    "users、password_reset_codes：账户、状态、认证版本和一次性重置码摘要。",
    "risk_events、risk_signals、risk_feedbacks：风险事件、证据及用户纠错。",
    "families、family_members、family_invitations、trusted_contacts：家庭关系与可信联系人。",
    "family_alerts、family_alert_recipients、family_alert_delivery_attempts、family_alert_policies：告警、接收人与投递审计。",
    "push_devices：设备安装标识、FCM Token与撤销状态。",
    "call_guard_reports：号码反馈HMAC指纹、脱敏值、类型、备注和复核状态。",
])
doc.add_heading("8.2 API分组", level=2)
add_table(["前缀", "主要能力"], [
    ("/api/v1/auth", "注册、登录、密码重置"),
    ("/api/v1/users/me", "当前用户、密码、推送设备"),
    ("/api/v1/risk", "文本、图片、语音、URL、事件、反馈、看板与报告"),
    ("/api/v1/call-guard", "号码分析、规则包、号码反馈、家庭求助"),
    ("/api/v1/families", "家庭、邀请、成员、联系人、告警与策略"),
    ("/api/v1/health", "服务健康检查"),
], [2100, 7260])
doc.add_paragraph("除健康检查、注册、登录和密码重置外，业务接口使用Authorization: Bearer <token>。精确字段、状态码和请求模型以运行时 /docs 或 /openapi.json 为准。", style="Small")

doc.add_heading("9. 安全与隐私设计", level=1)
add_table(["控制点", "当前实现", "说明"], [
    ("密码安全", "Argon2摘要", "不保存明文密码"),
    ("身份鉴权", "JWT + 认证版本", "修改密码后可使旧Token失效"),
    ("资源隔离", "用户所有权、家庭成员与角色校验", "降低越权和UUID枚举风险"),
    ("重置码", "限时、限重发、限失败次数、摘要保存", "减少验证码滥用"),
    ("上传安全", "文件类型、大小、音频时长校验；临时文件处理后清理", "减少恶意与无界上传"),
    ("号码隐私", "反馈保存HMAC-SHA256指纹和脱敏值", "在线号码分析仍需发送号码，不能宣称所有号码均不上传"),
    ("通话隐私", "来电护航不读取或录制实时通话音频", "RECORD_AUDIO权限用于独立语音检测，不等于通话录音"),
    ("结论边界", "风险提示、核验建议、误报/漏报反馈", "不构成司法、执法、金融或身份权威鉴定"),
], [1700, 3600, 4060])
callout("推荐宣传口径", "系统支持版本化规则同步和本地离线筛查；来电护航不录制实时通话内容，也不在本地保存通话历史。号码反馈采用HMAC指纹与脱敏值存储。分析结果用于风险提示和核验建议，而非权威鉴定。", "E9F7F5")

doc.add_heading("10. 测试与质量保障", level=1)
add_bullets([
    "后端测试按unit、integration、slow三层标记，覆盖快速逻辑、API/持久化和OCR/ASR/PDF等重依赖链路。",
    "Ruff执行代码规范与静态检查；mypy用于类型检查；pytest-cov支持覆盖率统计。",
    "Android使用JUnit及Compose/Espresso测试依赖，并提供testDebugUnitTest与compileDebugKotlin检查。",
    "数据库变更通过Alembic迁移管理，要求从空库顺序升级到最新版本。",
    "规则、API Schema和客户端模型的兼容性以OpenAPI及集成测试为判断依据。",
])
add_table(["检查项", "命令"], [
    ("后端规范", "uv run ruff check app tests"),
    ("后端单元", "uv run pytest -m unit"),
    ("后端集成", "uv run pytest -m integration"),
    ("后端慢速", "uv run pytest -m slow"),
    ("数据库迁移", "uv run alembic upgrade head"),
    ("Android", ".\\gradlew testDebugUnitTest compileDebugKotlin"),
], [2400, 6960])

doc.add_heading("11. 部署与运行", level=1)
add_numbered([
    "复制.env.example为.env，并替换数据库、MinIO和JWT等敏感配置。",
    "使用docker compose up -d postgres redis minio启动基础设施。",
    "进入apps/api_server，执行uv sync、Alembic迁移并启动Uvicorn。",
    "在Android Studio中打开apps/android_app，配置可被真机访问的API地址。",
    "如需FCM、邮件、短信、AI或语音模型，按对应指南配置凭证与本地模型。",
    "访问http://localhost:8000/docs检查OpenAPI，访问/api/v1/health检查服务状态。",
])
callout("配置安全", "不得提交真实.env、Firebase google-services.json、服务账号、模型密钥或生产凭证。生产环境应使用HTTPS、密钥管理、私网数据库与最小权限账号。", "FFF0F0", RED)

doc.add_heading("12. 项目创新点", level=1)
add_table(["创新方向", "项目体现"], [
    ("多模态统一识别", "文本、截图OCR、语音ASR、网址和来电号码进入统一风险表达与记录体系。"),
    ("规则+AI可解释融合", "以可审计规则为基础，以语义模型发现隐蔽行为链，并保留证据、置信度和版本。"),
    ("家庭协同干预", "从个人弹窗扩展到可信联系人、家庭告警、投递审计和处置闭环。"),
    ("风险前置", "把提醒放在转账、泄密、点击链接和通话求助等关键动作之前。"),
    ("隐私克制", "离线来电规则、号码反馈指纹化、不录制实时通话、非权威结论。"),
    ("工程完整度", "Android真机产品、后端API、数据库、迁移、通知、报告与测试形成可演示闭环。"),
], [2000, 7360])

doc.add_heading("13. 比赛实物演示建议", level=1)
add_table(["时间", "演示内容", "证明点"], [
    ("0:00-0:20", "真机首页、电脑服务端和项目定位", "真实可运行原型"),
    ("0:20-0:50", "诈骗文本检测：风险等级、证据、建议", "可解释识别"),
    ("0:50-1:15", "上传诈骗截图并展示OCR结果", "多模态能力"),
    ("1:15-1:35", "语音转写或网址检测", "跨渠道分析"),
    ("1:35-2:00", "手机A触发高风险，手机B收到家庭告警", "家庭协同核心创新"),
    ("2:00-2:20", "来电号码筛查或离线规则状态", "Android系统能力与隐私"),
    ("2:20-2:40", "历史记录、反馈和PDF报告", "闭环与可复盘"),
    ("2:40-2:55", "创新点与边界总结", "社会价值和可信表达"),
], [1350, 4650, 3360])

doc.add_heading("14. 当前边界与风险", level=1)
add_table(["类别", "当前边界", "建议"], [
    ("号码数据", "暂无运营商、公安或权威号码库实时接入", "宣传为本地维护规则与用户反馈复核，不宣称权威号码认定"),
    ("通话内容", "Android来电筛查不分析实时通话音频", "强调号码风险与通话求助，不宣传实时语音监听"),
    ("生产运维", "缺少完整分布式限流、审计平台、备份恢复与可观测性", "上线前补齐网关、监控、告警、备份和密钥托管"),
    ("基础设施", "Redis和MinIO已编排，但并非所有预期链路均已接入", "材料中区分“已部署基础设施”和“已在核心链路使用”"),
    ("测试", "Android端到端设备测试与真实PostgreSQL/通知通道验证仍需加强", "建立真机矩阵、回归用例和演示前检查表"),
    ("模型效果", "缺少公开基准、真实用户规模与持续评估流水线", "建设脱敏测试集，报告准确率、召回率、误报率及样本量"),
    ("运营后台", "号码复核与规则运营工具尚未形成完整产品", "增加审核台、规则发布、回滚、审计和权限管理"),
], [1500, 4500, 3360])

doc.add_heading("15. 未来持续完善重点", level=1)
doc.add_paragraph("以下内容属于后续建设方向，不代表当前版本已经完成。完善顺序以“先稳定、再验证、后扩展”为原则：首先保证比赛演示与核心闭环可靠，其次用真实脱敏数据验证效果，最后再推进生产化与生态合作。")
add_table(["优先级", "完善方向", "具体建设内容", "建议验收标准"], [
    ("P0", "演示稳定性", "提供一键演示账号与测试素材；预置文本、截图、语音、网址和号码案例；增加网络异常、服务不可用和权限缺失提示。", "完整演示连续运行10次无阻断；两部真机告警联动成功率达到100%。"),
    ("P0", "检测效果评估", "建设经授权、脱敏并分类标注的诈骗测试集，覆盖冒充公检法、刷单、投资、退款、验证码、屏幕共享等场景。", "公开样本量、准确率、召回率、误报率、漏报率及各类别指标，不只展示总体准确率。"),
    ("P0", "规则质量", "建立规则新增、复核、灰度发布、版本回滚和失效机制；补充否定语义、引用转发和反诈宣传文本等易误报场景。", "每条规则具有来源、版本、置信度、负责人、测试用例及回滚记录。"),
    ("P0", "家庭协同可靠性", "完善告警确认、未响应升级、重复告警合并、联系人值班顺序及离线补偿。", "告警可追踪到接收、查看、联系和处置状态；失败投递能够自动重试并解释原因。"),
    ("P1", "适老化体验", "增加大字体、高对比度、语音播报、一步求助、简化首页、方言友好提示和防误触确认。", "邀请老年用户完成可用性测试，统计任务完成率、平均耗时和误操作次数。"),
    ("P1", "OCR与语音增强", "改善低清截图、长截图、聊天气泡、多说话人、噪声、口音和方言识别；展示识别置信度并允许逐段纠正。", "在独立测试集上分别报告OCR字符准确率、ASR字错率及纠正前后风险判断变化。"),
    ("P1", "网址检测增强", "加入同形字域名、短链展开、重定向链、证书、域名年龄、页面仿冒特征和安全浏览数据源。", "限制跳转次数和请求时间；对不可访问网址给出明确的不确定性提示。"),
    ("P1", "号码风险能力", "建设号码审核后台，支持用户举报去重、证据复核、精确规则发布与申诉处理；探索合规的数据合作。", "未经人工或可信来源复核的举报不得进入自动拦截规则；所有变更可审计和回滚。"),
    ("P1", "AI模型治理", "建立提示词、模型、参数和融合策略版本管理；开展对抗提示、注入文本、偏见和幻觉测试。", "相同版本结果可复现；AI不可用时自动降级到规则分析；高危规则证据不可被模型静默覆盖。"),
    ("P1", "隐私与用户控制", "增加分场景授权、数据用途说明、保存期限、导出与删除入口；对日志、报告和通知内容实施更严格脱敏。", "形成数据清单与保留周期；用户可查看、导出和删除依法可处理的个人数据。"),
    ("P1", "安全加固", "接入HTTPS、API网关限流、密钥托管、依赖与镜像扫描、日志脱敏、异常登录保护和安全审计。", "完成威胁建模、自动扫描和基础渗透测试；高风险问题关闭后方可发布。"),
    ("P2", "运营管理后台", "建设用户、号码举报、规则、模型版本、告警投递和指标看板，并配置RBAC和审批流。", "关键操作要求权限校验、双人复核或审批，并保留不可抵赖的审计记录。"),
    ("P2", "生产可观测性", "补充结构化日志、指标、链路追踪、健康检查、告警、备份恢复和容量压测。", "定义SLO；完成数据库恢复演练、通知故障演练和峰值压测报告。"),
    ("P2", "端到端测试", "建立Android真机矩阵，覆盖不同厂商、Android版本、通知限制、弱网、离线规则和系统来电角色。", "核心用户旅程纳入自动化回归；每次发布生成兼容性与回归报告。"),
    ("P2", "多端与无障碍", "评估iOS受平台能力限制下的可实现范围；补充Web家属端、TalkBack语义和色弱友好设计。", "通过无障碍检查；多端数据权限和状态保持一致。"),
    ("P3", "试点与商业验证", "与社区养老、学校、银行或反诈教育机构开展小范围试点，验证使用频率、干预效果和服务模式。", "试点具备知情同意、退出机制、伦理与隐私审查，并形成可量化的社会价值报告。"),
], [800, 1450, 4410, 2700])

doc.add_heading("15.1 建议优先完成的比赛增强包", level=2)
add_bullets([
    "准备20-50条脱敏测试样本，形成基础准确率、召回率和误报案例说明。",
    "固定一套两部手机联动流程，并提供离线备用录屏，防止现场网络或推送通道异常。",
    "增加演示模式和一键清理测试数据，确保每次展示从相同状态开始。",
    "为每个创新点准备代码证据、界面证据、测试证据和边界说明各一项。",
    "制作一页竞品差异表，突出多模态、可解释证据、家庭协同和隐私克制。",
])

doc.add_heading("15.2 建议建立的量化指标", level=2)
add_table(["指标类别", "核心指标", "意义"], [
    ("识别效果", "准确率、召回率、精确率、F1、误报率、漏报率", "证明风险识别不是只靠个别演示案例"),
    ("多模态质量", "OCR字符准确率、ASR字错率、URL解析成功率", "定位风险分析前置转换环节的问题"),
    ("家庭协同", "告警送达率、查看率、响应时间、处置完成率", "证明家庭干预闭环是否真正有效"),
    ("来电护航", "本地判断耗时、规则命中率、同步成功率、误拦截率", "验证系统响应窗口与安全边界"),
    ("稳定性", "崩溃率、API错误率、P95延迟、通知重试成功率", "衡量产品是否具备持续使用能力"),
    ("用户体验", "任务完成率、平均完成时间、求助步骤数、满意度", "验证老年人与普通家庭成员是否易用"),
], [1700, 4100, 3560])

doc.add_heading("16. 后续路线图", level=1)
add_table(["阶段", "目标", "建议交付物"], [
    ("第一阶段：比赛增强", "稳定真机演示、补充测试数据与答辩材料", "一键演示数据、两机联动、3分钟视频、指标表"),
    ("第二阶段：小范围试点", "家庭/社区场景验证可用性与误报体验", "知情同意、脱敏样本、用户访谈、误报处理流程"),
    ("第三阶段：运营能力", "号码审核、规则发布、版本回滚和审计", "管理后台、RBAC、审批流、规则质量指标"),
    ("第四阶段：生产化", "安全、可靠性、监控、备份和合规", "HTTPS、限流、日志脱敏、告警、灾备、渗透测试"),
    ("第五阶段：生态合作", "探索银行、社区养老、运营商及反诈教育合作", "数据合作边界、接口规范、联合试点与商业模式"),
], [1900, 3600, 3860])

doc.add_heading("附录A：对外介绍参考", level=1)
doc.add_paragraph("真信盾是一套面向家庭场景的多模态智能反诈与可信联系系统，可对聊天文本、短信截图、语音文件、可疑网址和陌生来电号码进行风险分析，并给出风险等级、命中证据及行动建议。项目将可解释规则与可选AI语义复核结合，把个人风险提醒进一步延伸到可信联系人和家庭告警，在不录制实时通话内容的前提下，形成“识别风险—解释证据—家人协同—反馈复盘”的防护闭环。")
doc.add_page_break()
doc.add_heading("附录B：准确表述与禁用表述", level=1)
add_table(["推荐表述", "避免表述"], [
    ("风险提示、核验建议", "权威鉴定、100%识别"),
    ("本地维护的版本化号码规则", "运营商/公安权威号码库"),
    ("来电号码风险筛查", "实时监听或分析通话内容"),
    ("确认风险号码可选拦截", "所有骚扰电话自动拦截"),
    ("规则与AI融合分析", "AI能够准确判断所有诈骗"),
    ("面向家庭协同的产品原型", "已大规模商业化部署"),
], [4680, 4680])

# Keep tables readable and mark headers repeated.
for table in doc.tables:
    if table.rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"), "true"); tr_pr.append(repeat)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)

doc.core_properties.title = "真信盾项目功能与技术说明书"
doc.core_properties.subject = "项目功能、技术架构、安全设计与比赛材料"
doc.core_properties.author = "True Shield Project Team"
doc.core_properties.keywords = "真信盾, 反诈, Android, FastAPI, 家庭协同, 多模态"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
